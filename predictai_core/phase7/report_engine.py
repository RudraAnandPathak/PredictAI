"""Single report content model with PDF, DOCX, and PPTX exporters."""
from __future__ import annotations
from datetime import datetime
from io import BytesIO
from typing import Any
import json
import re


def build_report(title: str, sections: list[dict[str, Any]], metadata: dict[str, Any] | None = None) -> dict[str, Any]:
    clean = []
    for section in sections:
        heading = str(section.get('heading', 'Section')).strip() or 'Section'
        body = section.get('body', '')
        if isinstance(body, (dict, list, tuple)):
            body = json.dumps(body, ensure_ascii=False, indent=2, default=str)
        clean.append({'heading': heading, 'body': str(body).strip()})
    return {'title': str(title or 'PredictAI Report').strip() or 'PredictAI Report', 'created_at': datetime.now().astimezone().isoformat(timespec='seconds'), 'metadata': metadata or {}, 'sections': clean}


def quality_check(report: dict[str, Any]) -> dict[str, Any]:
    issues = []
    if not report.get('title','').strip(): issues.append('Report title is missing.')
    sections = report.get('sections', [])
    if not sections: issues.append('No report sections were added.')
    for section in sections:
        if not str(section.get('body','')).strip(): issues.append(f"Section '{section.get('heading','Untitled')}' is empty.")
    joined = '\n'.join(str(s.get('body','')) for s in sections).lower()
    if any(word in joined for word in ('prediction','model','forecast')) and not any(word in joined for word in ('limitation','caution','uncertainty','not causal','not establish causation')):
        issues.append('Consider adding a model limitation or uncertainty statement.')
    return {'passed': not issues, 'issues': issues, 'checked_sections': len(sections), 'note': 'Automated checks catch common omissions; they do not certify factual correctness.'}


def export_txt(report: dict[str, Any]) -> bytes:
    parts = [report['title'], f"Generated: {report.get('created_at','')}", '']
    for section in report.get('sections', []): parts += [section['heading'], '-' * len(section['heading']), section['body'], '']
    return '\n'.join(parts).encode('utf-8')


def export_pdf(report: dict[str, Any]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted, PageBreak
    out = BytesIO(); doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
    styles = getSampleStyleSheet(); styles.add(ParagraphStyle(name='ReportTitle', parent=styles['Title'], alignment=TA_LEFT, textColor=colors.HexColor('#4338CA')))
    story = [Paragraph(_escape(report['title']), styles['ReportTitle']), Spacer(1, 6*mm), Paragraph('Generated: '+_escape(report.get('created_at','')), styles['Normal']), Spacer(1, 5*mm)]
    for section in report.get('sections', []):
        story += [Paragraph(_escape(section['heading']), styles['Heading2']), Spacer(1, 2*mm), Preformatted(section['body'][:50000], styles['Code']), Spacer(1, 4*mm)]
    doc.build(story); return out.getvalue()


def _escape(value: str) -> str:
    return str(value).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')


def export_docx(report: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document(); doc.add_heading(report['title'], 0)
    doc.add_paragraph('Generated: '+report.get('created_at',''))
    for section in report.get('sections', []):
        doc.add_heading(section['heading'], level=1)
        for block in re.split(r'\n\s*\n', section['body']):
            if block.strip(): doc.add_paragraph(block.strip())
    out = BytesIO(); doc.save(out); return out.getvalue()


def export_pptx(report: dict[str, Any]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text = report['title']; slide.placeholders[1].text = 'Generated: '+report.get('created_at','')
    for section in report.get('sections', []):
        slide = prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text = section['heading']
        text = section['body']
        # Keep each slide readable; long sections are split into continuation slides.
        chunks = [text[i:i+1200] for i in range(0, max(1,len(text)), 1200)] or ['']
        for index, chunk in enumerate(chunks):
            if index:
                slide = prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text = section['heading']+' (continued)'
            slide.placeholders[1].text = chunk[:1200]
            for p in slide.placeholders[1].text_frame.paragraphs: p.font.size = Pt(16)
    out = BytesIO(); prs.save(out); return out.getvalue()
