from pathlib import Path
import pandas as pd
import pytest
from docx import Document
from document_reader import read_document


def test_unsupported_extension_is_clear(tmp_path: Path):
    p = tmp_path / "x.json"; p.write_text("{}", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported file format"):
        read_document(p)


def test_missing_file_is_clear(tmp_path: Path):
    with pytest.raises(FileNotFoundError): read_document(tmp_path / "missing.csv")


def test_csv_semicolon_and_bom(tmp_path: Path):
    p = tmp_path / "data.csv"; p.write_text("\ufeffname;score\nAsha;10\nRavi;20\n", encoding="utf-8")
    result = read_document(p)
    assert result["type"] == "tabular" and isinstance(result["content"], pd.DataFrame)
    assert result["content"].shape == (2, 2)
    assert list(result["content"].columns) == ["name", "score"]


def test_text_and_markdown(tmp_path: Path):
    for suffix, label in ((".txt", "TXT"), (".md", "MD")):
        p = tmp_path / f"notes{suffix}"; p.write_text("Hello PredictAI", encoding="utf-8")
        result = read_document(p)
        assert result["type"] == "text" and result["format"] == label
        assert result["content"] == "Hello PredictAI"


def test_rtf_extracts_basic_text(tmp_path: Path):
    p = tmp_path / "notes.rtf"; p.write_text(r"{\rtf1\ansi Hello\par PredictAI}", encoding="utf-8")
    result = read_document(p)
    assert result["type"] == "text" and "Hello" in result["content"] and "PredictAI" in result["content"]


def test_docx_includes_table_text(tmp_path: Path):
    p = tmp_path / "table.docx"
    doc = Document(); doc.add_paragraph("Overview")
    table = doc.add_table(rows=2, cols=2); table.cell(0,0).text="Name"; table.cell(0,1).text="Score"
    table.cell(1,0).text="Asha"; table.cell(1,1).text="95"; doc.save(p)
    result = read_document(p)
    assert result["type"] == "docx" and result["table_count"] == 1
    assert any("Asha" in line and "95" in line for line in result["content"])


def test_xlsx_first_sheet_compatible(tmp_path: Path):
    p = tmp_path / "data.xlsx"
    pd.DataFrame({"name": ["Asha"], "score": [10]}).to_excel(p, index=False)
    result = read_document(p)
    assert result["type"] == "tabular" and result["content"].shape == (1, 2)
