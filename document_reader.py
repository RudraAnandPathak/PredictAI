"""Backward-compatible document extraction for PredictAI Phase 2.

PDF returns page text and optional best-effort OCR metadata; DOCX returns paragraph
and table text; CSV/XLSX/XLS/ODS return a pandas DataFrame (first worksheet for
now); TXT/Markdown/RTF return extracted text. OCR requires a system Tesseract
installation and is best-effort: lack of OCR support never prevents normal PDFs
from being read.
"""
from __future__ import annotations
import csv
import html
import re
from pathlib import Path
from typing import Any
import pandas as pd
import pymupdf
from docx import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".csv", ".xlsx", ".xls", ".ods", ".txt", ".md", ".rtf"}


def _read_csv(path: Path) -> pd.DataFrame:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return pd.read_csv(path, encoding=encoding, sep=None, engine="python", on_bad_lines="error")
        except (UnicodeDecodeError, pd.errors.ParserError, csv.Error) as exc:
            last_error = exc
    raise ValueError("Could not read this CSV file. Check that it is valid and has a consistent number of fields per row.") from last_error


def _read_pdf(path: Path, enable_ocr: bool = True) -> dict[str, Any]:
    pages: list[dict[str, Any]] = []
    ocr_pages: list[int] = []
    warnings: list[str] = []
    tables: list[dict[str, Any]] = []
    try:
        with pymupdf.open(path) as pdf:
            if pdf.is_encrypted:
                raise ValueError("This PDF is encrypted or password-protected.")
            for page_number, page in enumerate(pdf, start=1):
                text = (page.get_text("text") or "").strip()
                # Only attempt OCR for pages with little/no selectable text.
                if enable_ocr and len(text) < 30:
                    try:
                        text_page = page.get_textpage_ocr(language="eng", dpi=200, full=True)
                        ocr_text = (page.get_text("text", textpage=text_page) or "").strip()
                        if len(ocr_text) > len(text):
                            text = ocr_text
                            if text:
                                ocr_pages.append(page_number)
                    except Exception as exc:
                        # Do not fail document extraction if the optional Tesseract OCR dependency is absent.
                        if not warnings:
                            warnings.append("Automatic OCR was unavailable. Install Tesseract OCR and ensure it is on PATH to extract scanned PDF pages.")
                pages.append({"page": page_number, "text": text})
                # Best-effort extraction of real PDF tables. Complex layouts and image tables may not be recognized.
                try:
                    finder = page.find_tables()
                    for table_number, table in enumerate(finder.tables, start=1):
                        raw_rows = table.extract() or []
                        cleaned = [[("" if cell is None else str(cell).strip()) for cell in row] for row in raw_rows]
                        cleaned = [row for row in cleaned if any(cell for cell in row)]
                        if len(cleaned) < 2:
                            continue
                        headers = cleaned[0]
                        if not any(headers) or len(set(headers)) != len(headers):
                            headers = [f"Column {i + 1}" for i in range(len(cleaned[0]))]
                            body = cleaned
                        else:
                            headers = [header or f"Column {i + 1}" for i, header in enumerate(headers)]
                            body = cleaned[1:]
                        width = len(headers)
                        body = [(row + [""] * width)[:width] for row in body]
                        frame = pd.DataFrame(body, columns=headers)
                        # Convert columns that are mostly numeric so statistics and plots work immediately.
                        for column in frame.columns:
                            values = frame[column].astype(str).str.replace(",", "", regex=False).str.replace("%", "", regex=False).str.strip()
                            numeric = pd.to_numeric(values, errors="coerce")
                            nonempty = values.ne("").sum()
                            if nonempty and numeric.notna().sum() / nonempty >= 0.7:
                                frame[column] = numeric
                        tables.append({"page": page_number, "table_number": table_number, "dataframe": frame})
                except Exception:
                    # Table extraction must never prevent normal PDF text extraction.
                    pass
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not open this PDF. It may be damaged or invalid.") from exc
    return {"type": "pdf", "content": pages, "ocr_pages": ocr_pages, "warnings": warnings, "tables": tables}


def _read_docx(path: Path) -> dict[str, Any]:
    try:
        document = Document(path)
    except Exception as exc:
        raise ValueError("Could not open this Word document. It may be damaged or invalid.") from exc
    content = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    table_count = len(document.tables)
    for index, table in enumerate(document.tables, start=1):
        content.append(f"[Table {index}]")
        for row in table.rows:
            cells = [cell.text.replace("\\n", " / ").strip() for cell in row.cells]
            if any(cells):
                content.append(" | ".join(cells))
    return {"type": "docx", "content": content, "table_count": table_count}


def _read_rtf(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    # Lightweight RTF text extraction for common documents. It is not a full RTF renderer.
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", lambda m: bytes([int(m.group(0)[2:], 16)]).decode("cp1252", errors="replace"), raw)
    raw = re.sub(r"\\u(-?\d+)\??", lambda m: chr(int(m.group(1)) % 65536), raw)
    raw = re.sub(r"\\(par|line)\b", "\n", raw)
    raw = re.sub(r"\\tab\b", "\t", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", raw)
    raw = raw.replace("{", "").replace("}", "")
    raw = raw.replace("\\\\", "\\").replace("\\{", "{").replace("\\}", "}")
    return html.unescape(raw).strip()


def read_document(file_path: str | Path) -> dict[str, Any]:
    """Extract a supported file while retaining app.py's type/content contract."""
    path = Path(file_path)
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"The selected file does not exist: {path.name}")
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file format: {ext or '(no extension)'}. Supported formats: {supported}")
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext in {".txt", ".md"}:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = path.read_text(encoding="cp1252", errors="replace")
        return {"type": "text", "format": ext[1:].upper(), "content": text}
    if ext == ".rtf":
        return {"type": "text", "format": "RTF", "content": _read_rtf(path)}
    if ext == ".csv":
        return {"type": "tabular", "content": _read_csv(path)}
    engine = {".xlsx": "openpyxl", ".xls": "xlrd", ".ods": "odf"}[ext]
    try:
        dataframe = pd.read_excel(path, sheet_name=0, engine=engine)
    except ImportError as exc:
        package = {".xlsx": "openpyxl", ".xls": "xlrd", ".ods": "odfpy"}[ext]
        raise RuntimeError(f"Reading {ext} files requires the optional '{package}' package. Install the updated requirements.txt.") from exc
    except Exception as exc:
        raise ValueError(f"Could not read this spreadsheet. Check that it is a valid {ext} file and that its dependencies are installed.") from exc
    return {"type": "tabular", "content": dataframe}


if __name__ == "__main__":
    print("PredictAI document reader is ready.")
